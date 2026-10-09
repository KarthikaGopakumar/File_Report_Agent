from fastapi.responses import FileResponse
from fastapi import FastAPI, UploadFile, File
from graph_excel import graph
from docx import Document
from fastapi import HTTPException

app = FastAPI()

def make_docx(report_text):
    doc=Document()    
    doc.add_heading("Sales Report",level=1)
    doc.add_paragraph(report_text)
    doc.save("report.docx")

@app.post("/report")
async def report(file: UploadFile = File(...)):
    if not file.filename.endswith("xlsx"):
         raise HTTPException(status_code=400, detail="Please upload an .xlsx file.")
    contents = await file.read()
    with open("uploaded.xlsx", "wb") as f:
        f.write(contents)
    result=graph.invoke({"file_path": "uploaded.xlsx"}) 
    if result.get("error"):
        raise HTTPException(status_code=400, detail=result["error"])
    make_docx(result["report"])   
    return FileResponse("report.docx", filename="report.docx")